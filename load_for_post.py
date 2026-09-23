#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
load_for_post.py — выгрузка утверждённых документов с lks.dap.gov.ru

Чистый Python (стандартная библиотека: urllib, http.cookiejar) — без браузера.
Вход: POST /internal/auth/authenticator/api/internalauth/auth (Authorization: Basic)
Данные: POST ApprovalDocument/List + GET ApprovalDocument/GetFinalPdf?id=<id>
       (проверено: GetFinalPdf = вкладка «Печатная форма» — финальный PDF
       со штампом об электронной подписи; GetPdf — черновой вариант без штампа)

Использование:
    python3 load_for_post.py [ДД.ММ.ГГГГ] [--post-only | --no-post]
(без аргумента — сегодняшняя дата)

Стадия 1 — выгрузка утверждённых PDF (Извещение/Протокол/Определение/Постановление)
          в out/<ДД.ММ.ГГГГ>/<тип>/.
Стадия 2 — почта РФ: получатель определяется из текста каждого PDF (ФИО/название,
          ИНН, адрес); документы разных типов на одного получателя (ключ — ИНН,
          резерво — нормализованное ФИО/название) сливаются в один PDF;
          заполняется реестр registry.xlsx (по шаблону registry-template.xlsx,
          тип письма = 1, КПП и доп. услуга ЮЗЭУВ не заполняются);
          итоговые архивы out/<ДД.ММ.ГГГГ>/post.zip = сжатые PDF + registry.xlsx.
          Лимит: не более ZIP_MAX_DOCS (50) писем в одном архиве; при превышении
          архивы разбиваются: post1.zip, post2.zip, ... — каждый со своим
          registry.xlsx (только свои получатели).
          Для стадии 2 нужны пакеты: pypdf, pdfplumber, openpyxl.
"""

import base64
import datetime
import json
import os
import re
import shutil
import sys
import urllib.error
import urllib.request
import zipfile
import http.cookiejar

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")
OUT_ROOT = os.path.join(BASE_DIR, "out")

HOST = "https://lks.dap.gov.ru"
AUTH_URL = HOST + "/internal/auth/authenticator/api/internalauth/auth?loaderKey=default"
LIST_URL = HOST + "/public/rosstat/ouzd/web_api/api/ApprovalDocument/List"
ZIP_MAX_DOCS = 50  # максимум писем (получателей) в одном post*.zip
PDF_URL = HOST + "/public/rosstat/ouzd/web_api/api/ApprovalDocument/GetFinalPdf?id={id}"
XML_URL = HOST + "/public/rosstat/ouzd/web_api/api/ApprovalDocument/GetXml?id={id}"
LOGIN_PAGE = HOST + "/internal/authutil/auth/login"

TYPE_IDS = [
    ("Извещение", 48),
    ("Протокол", 36),
    ("Определение", 38),
    ("Постановление", 37),
]

SUBSYSTEM_ID = 80
USER_AGENT = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/152.0.0.0 Safari/537.36"


def load_config():
    if not os.path.exists(CONFIG_PATH):
        sys.exit("ОШИБКА: нет config.json: %s\n"
                 "Создайте файл вида {\"login\": \"...\", \"password\": \"...\"}"
                 % CONFIG_PATH)
    with open(CONFIG_PATH) as f:
        cfg = json.load(f)
    if "login" not in cfg or "password" not in cfg:
        sys.exit("ОШИБКА: в config.json должны быть поля login и password")
    return cfg


def make_opener(login_name, password):
    jar = http.cookiejar.CookieJar()
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    token = base64.b64encode(("%s:%s" % (login_name, password)).encode()).decode()
    op.addheaders = [
        ("User-Agent", USER_AGENT),
        ("Authorization", "Basic " + token),
    ]
    return op


def login(op, login_name, password, timeout=60):
    """Открыть страницу логина (сессионный cookie) и выполнить Basic-авторизацию."""
    req = urllib.request.Request(LOGIN_PAGE, method="GET")
    op.open(req, timeout=timeout).read()

    req = urllib.request.Request(AUTH_URL, data=b"", method="POST")
    resp = op.open(req, timeout=timeout)
    body = resp.read().decode(errors="replace")
    if resp.status not in (200, 201):
        sys.exit("ОШИБКА авторизации: HTTP %s\n%s" % (resp.status, body[:500]))
    try:
        data = json.loads(body)
        if data.get("userAccountStatus") not in (None, "ACTIVE"):
            sys.exit("ОШИБКА авторизации: статус учётной записи %s" % data.get("userAccountStatus"))
    except (ValueError, TypeError):
        pass  # допустимо — главное, что HTTP 200 и cookie получены
    print("Вход выполнен.")


def list_documents(op, date_iso, type_id, timeout=90):
    """Список утверждённых документов заданной даты для типа."""
    year = int(date_iso[:4])
    body = {
        "page": 0,
        "start": 0,
        "limit": 500,
        "dataFilter": {
            "Group": 2,
            "Filters": [
                {"DataIndex": "State", "Value": "Утверждено", "Operand": 6, "Type": "text"},
                {"DataIndex": "SignDate", "Value": date_iso, "Operand": 6, "Type": "date"},
            ],
        },
        "sort": [],
        "sorters": [],
        "accessFilter": 0,
        "documentsFilter": 2,
        "receiptYear": year,
        "subSystemId": SUBSYSTEM_ID,
        "documentTypeIds": [type_id],
    }
    req = urllib.request.Request(LIST_URL, data=json.dumps(body).encode(), method="POST")
    req.add_header("Content-Type", "application/json")
    resp = op.open(req, timeout=timeout)
    data = json.loads(resp.read().decode())
    return data.get("data", []), data.get("totalCount", 0)


def get_pdf(op, doc_id, timeout=180):
    """Скачать PDF документа; вернуть байты."""
    req = urllib.request.Request(PDF_URL.format(id=doc_id), method="GET")
    resp = op.open(req, timeout=timeout)
    raw = resp.read()
    text = raw.decode(errors="replace").strip()
    if text.startswith('"'):
        text = json.loads(raw)
    if text.startswith("%PDF"):
        return text.encode()
    pdf = base64.b64decode(text)
    if not pdf.startswith(b"%PDF"):
        raise ValueError("Ответ не похож на PDF (начало: %r)" % pdf[:16])
    return pdf


def get_xml(op, doc_id, timeout=90):
    """Форма (XML) утверждённого документа: JSON-ответ содержит base64-строку XML."""
    req = urllib.request.Request(XML_URL.format(id=doc_id), method="GET")
    resp = op.open(req, timeout=timeout)
    raw = resp.read().decode(errors="replace").strip()
    if raw.startswith('"'):
        raw = json.loads(raw)
    if raw.startswith("<?xml"):
        return raw.encode()
    data = base64.b64decode(raw)
    if not data.startswith(b"\xef\xbb\xbf") and not data.lstrip().startswith(b"<"):
        raise ValueError("Ответ не похож на XML (начало: %r)" % data[:16])
    return data


CONFIRM_HEAD_LIMIT = 500


def _is_unwanted_text(txt):
    """Причина, ПОЧЕМУ документ не выгружать (или None — обычный документ).
    Оба исключения — «обёрточные» документы, узнаются по ЗАГОЛОВКУ (первые ~450–600 зн.):
      - «письмо-уведомление о документе» — обложка письма о раннем документе;
      - «подтверждение вручения письма-уведомления» — но если заголовок —
        настоящий ПРОТОКОЛ/ОПРЕДЕЛЕНИЕ/ПОСТАНОВЛЕНИЕ/ИЗВЕЩЕНИЕ №, не трогаем
        (в реальных документах фраза про подтверждение есть в перечне приложений)."""
    head = re.sub(r"\s+", " ", txt or "").strip()
    if not head:
        return None
    if re.search(r"письмо[\s-]+уведомлени\w+\s+о\s+документ\w*", head[:450], re.I):
        return "письмо-уведомление о документе"
    if re.search(r"(ПРОТОКОЛ|ОПРЕДЕЛЕНИЕ|ПОСТАНОВЛЕНИЕ|ИЗВЕЩЕНИЕ)\s*№", head[:600], re.I):
        return None
    if re.search(r"подтверждени\w*\s+вручени", head[:CONFIRM_HEAD_LIMIT], re.I):
        return "подтверждение вручения письма-уведомления"
    return None


def _is_unwanted_pdf(pdf_bytes):
    """То же, по бинарному PDF (первая страница)."""
    if not pdf_bytes or len(pdf_bytes) < 2048:
        return None
    try:
        import io as _io
        from pypdf import PdfReader
        r = PdfReader(_io.BytesIO(pdf_bytes))
        txt = r.pages[0].extract_text() or ""
    except Exception:
        return None
    return _is_unwanted_text(txt)


def _normalize_yo_pdf(pdf_path):
    """Заменяет в тексте PDF большую букву Ё (U+0401) на Е (U+0415).

    Почта РФ отклоняет письма, в которых ФИО содержит букву Ё
    (в почтовых ФИО по стандарту пишется Е). Заменяется только
    большая Ё (код 0x9C в подмножестве шрифта); малое ё (0xBC)
    проходит валидацию — не трогаем. Возвращает число замен.
    """
    try:
        import pikepdf
    except ImportError:
        print("  ВНИМАНИЕ: нет pikepdf — не могу нормализовать Ё (pip install pikepdf)")
        return 0
    n = 0
    try:
        pdf = pikepdf.open(pdf_path)
        for page in pdf.pages:
            cont = page.Contents
            streams = list(cont) if isinstance(cont, pikepdf.Array) else [cont]
            for s in streams:
                raw = s.read_bytes()
                # код 0x9C встречается как восьричный эскейп \234 в content-потоке
                c = raw.count(b"\x5c234") + raw.count(b"\x9c")
                if c:
                    new = raw.replace(b"\x5c234", b"\x5c305").replace(b"\x9c", b"\xC5")
                    s.write(new)
                    n += c
        if n:
            tmp = pdf_path + ".norm.tmp"
            pdf.save(tmp)
            pdf.close()
            os.replace(tmp, pdf_path)
        else:
            pdf.close()
    except Exception as e:
        print("  ВНИМАНИЕ: не удалось нормализовать Ё в %s: %s" % (os.path.basename(pdf_path), e))
        return 0
    return n


def sanitize(name):
    return (name or "").replace("/", "_").replace("\\", "_").strip() or "doc"


def ask_date(spec=None):
    if spec is not None:
        s = spec.strip()
    elif len(sys.argv) > 1:
        s = sys.argv[1].strip()
    else:
        try:
            s = input("Дата утверждения (ДД.ММ.ГГГГ, по умолчанию сегодня): ").strip()
        except EOFError:
            s = ""
    if not s:
        d = datetime.date.today()
        return d
    else:
        # форматы: ДД.ММ.ГГГГ / ГГГГ-ММ-ДД / ГГГГММДД
        s = s.strip()
        d = None
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%Y.%m.%d", "%d/%m/%Y"):
            try:
                d = datetime.datetime.strptime(s, fmt).date()
                break
            except ValueError:
                pass
        if d is None and len(s) == 8 and s.isdigit():
            d = datetime.date(int(s[:4]), int(s[4:6]), int(s[6:8]))
        if d is None:
            sys.exit("ОШИБКА: не удалось распознать дату: %r" % s)
        return d


def do_download(cfg, d, date_iso, date_ddmmyyyy):
    op = make_opener(cfg["login"], cfg["password"])
    login(op, cfg["login"], cfg["password"])

    results = []
    for type_name, type_id in TYPE_IDS:
        try:
            docs, total = list_documents(op, date_iso, type_id)
        except urllib.error.HTTPError as e:
            print("  %-16s ОШИБКА списка: HTTP %s" % (type_name, e.code))
            continue
        print("%s: %d" % (type_name, total))

        type_dir = os.path.join(OUT_ROOT, date_ddmmyyyy, type_name)
        os.makedirs(type_dir, exist_ok=True)

        ok = 0
        for doc in docs:
            stem = sanitize(doc.get("DocumentNumber"))
            fname = stem + ".pdf"
            fpath = os.path.join(type_dir, fname)
            xml_path = os.path.join(type_dir, stem + ".xml")
            need_pdf = not (os.path.exists(fpath) and os.path.getsize(fpath) > 0)
            need_xml = not (os.path.exists(xml_path) and os.path.getsize(xml_path) > 0)
            if not (need_pdf or need_xml):
                ok += 1
                print("  [уже есть] %s" % fname)
                continue
            if need_pdf:
                try:
                    pdf = get_pdf(op, doc["Id"])
                    why = _is_unwanted_pdf(pdf)
                    if why:
                        print("  [пропущено] %s — %s, не сохраняется" % (fname, why))
                        continue
                except Exception as e:
                    print("  [ошибка] %s: %s" % (fname, e))
                    continue
                tmp = fpath + ".part"
                with open(tmp, "wb") as f:
                    f.write(pdf)
                os.replace(tmp, fpath)
                print("  [%6d байт] %s" % (len(pdf), fname))
            if need_xml:
                try:
                    xml = get_xml(op, doc["Id"])
                    tmpx = xml_path + ".part"
                    with open(tmpx, "wb") as f:
                        f.write(xml)
                    os.replace(tmpx, xml_path)
                    print("  [%6d байт] %s" % (len(xml), stem + ".xml"))
                except Exception as e:
                    print("  [ошибка XML-формы] %s: %s" % (fname, e))
            ok += 1
        results.append((type_name, total, ok))

    print("\nИтог:")
    tot_listed = tot_saved = 0
    for type_name, total, ok in results:
        tot_listed += total
        tot_saved += ok
        print("  %-16s найдено: %d, сохранено/готово: %d" % (type_name, total, ok))
    print("  Всего документов: %d. Каталог: %s" % (tot_listed, os.path.join(OUT_ROOT, date_ddmmyyyy)))
    return 0 if tot_listed == tot_saved or tot_listed == 0 else 1


# ============================================================
# Стадия 2 — почта РФ: получатели из PDF, слияние, реестр, post.zip
# ============================================================

TYPE_ORDER = {"Извещение": 0, "Протокол": 1, "Определение": 2, "Постановление": 3}

_NAME_STOP = {
    "правонарушитель", "правонарушителя", "субъект", "состава", "составы", "состав",
    "статье", "части", "предусмотренного", "предусмотренной", "частью", "коап", "рф", "повторное",
    "совершение", "административного", "правонарушения",

    "индивидуальный", "индивидуальная", "индивидуальное", "индивидуального",
    "индивидуальной", "предприниматель", "предпринимателя", "предпринимателей",
    "гражданин", "гражданка", "гражданина", "лицо", "лица", "наименование",
    "организация", "организации", "нарушитель", "нарушителя", "совершившее",
    "совершивший", "совершившая", "в", "с", "и", "или", "адресат", "нарушитель-", 
}


def _flat_text(t):
    """Убирает переводы строк и дефисные переводы слов."""
    t = t.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"[-–] ?\n", "", t)
    t = t.replace("\n", " ")
    t = re.sub(r"[ \t]+", " ", t)
    return t.strip()


def _clean_name_tokens(seg):
    """Токены ФИО/наименования: выкидывает служебные слова, склеенные нижним регистром."""
    out = []
    for part in re.split(r"[\s\-]+", seg):
        if not part:
            continue
        if not re.search(r"[\u0400-\u04FF]", part):
            continue
        if part.lower() in _NAME_STOP:
            continue
        if len(part) < 2:
            continue
        # склеенное нижнем регистром служебное слово:
        # "индивидуальныйпредпринимательБУБНОВАНИНА" -> "БУБНОВАНИНА"
        mm = re.match(r"^[а-яёa-z]+([А-ЯЁA-Z][\u0400-\u04FFA-Za-z \-]*)$", part)
        if mm:
            part = mm.group(1)
            if part.lower() in _NAME_STOP or len(part) < 2:
                continue
        out.append(part)
    fin = []
    for tok in out:
        sp = _split_glued_caps(tok)
        fin.extend(sp if sp else [tok])
    return fin


_PATRONYMIC_SUFFIXES = ("ОВИЧ", "ЕВИЧ", "ОВНА", "ЕВНА", "ИЧНА", "ИВНА")

_FIRST_NAMES = {
 "АЛЕКСАНДР","АЛЕКСЕЙ","АНДРЕЙ","АНТОН","АРОН","АРСЕНИЙ","АРСЕН","АРТЁМ","АРТЕМ",
 "БОРИС","ВАДИМ","ВАЛЕНТИН","ВАЛЕНТИНА","ВАЛЕРИЙ","ВАСИЛИЙ","ВЕРА","ВЕНИАМИН",
 "ВИКТОР","ВИКТОРИЯ","ВИТАЛИЙ","ВЛАДИМИР","ГЕННАДИЙ","ГЕОРГИЙ","ГРИГОРИЙ",
 "ДАНИИЛ","ДАРЬЯ","ДЕНИС","ДМИТРИЙ","ЕГОР","ЕВГЕНИЙ","ЕВГЕНИЯ","ЕЛЕНА","ЕКАТЕРИНА",
 "ЗИНАИДА","ИВАН","ИГОРЬ","ИОСИФ","ИРИНА","ИСКРА","КИРИЛЛ","КОНСТАНТИН","КСЕНИЯ",
 "ЛАРИСА","ЛЕОНИД","ЛЮДМИЛА","МАКСИМ","МАРИЯ","МАРИНА","МИХАИЛ","НАДЕЖДА","НАТАЛЬЯ",
 "НАТАЛИЯ","НИКИТА","НИКОЛАЙ","НИНА","ОЛЕГ","ОЛЬГА","ПАВЕЛ","ПЕТР","ПЁТР","РАИСА",
 "РОБЕРТ","РОМАН","РУСЛАН","СВЕТЛАНА","СЕРГЕЙ","СОФЬЯ","ТАИСИЯ","ТАТЬЯНА","ТАМАРА",
 "ОЛЕНА","МИХАЙЛ","АНА","ЕКА",
 "ТИМУР","ФЕДОР","ФЕЛИКС","ФИЛИПП","ЮЛИЯ","ЮРИЙ","ЯРОСЛАВ","АГАФЬЯ","ЖАННА","ЗОЯ","ДМИТРИ","АНДРЕ","ИГОР","ЕВГЕНИ","АНАТОЛИ","ГЕННАДИ","АРКАДИ","СТАНИСЛА",
}


def _split_glued_caps(tok):
    """Склеенные ПРОПИСНЫМИ фамилия+имя[+отчество] -> список слов.
    1) хвост в отчестве (ОВИЧ/ЕВИЧ/ОВНА/ЕВНА/ИЧНА/ИВНА): длина отчества такая,
       при которой голова выглядит как фамилия и ствол отчества — имя/ствол;
    2) имя по словарю (с хвоста), голова должна выглядеть как фамилия;
    3) эвристика: фамилия оканчивается ОВ/ЕВ, имя >= 4 букв
       (не применяется к словам, оканчивающимся как отчество).
    Не удалось — None."""
    t = tok
    if not (len(t) >= 8 and re.fullmatch(r"[\u0410-\u042F]+", t)):
        return None
    _FAM_ENDS = ("ОВ", "ЕВ", "ОВА", "ЕВА", "ИНА")
    _PAT_ENDS = ("ОВИЧ", "ЕВИЧ", "ОВНА", "ЕВНА", "ИЧНА", "ИВНА")

    def _ok(head):
        if len(head) < 5:
            return False
        if head.endswith(_FAM_ENDS):
            return True
        for i in range(5, len(head) - 3):
            if head[:i].endswith(("ОВ", "ЕВ")) and len(head) - i >= 4:
                return True
        return False

    parts = []
    if t.endswith(_PAT_ENDS):
        best = None
        # длина отчества до 15 (АЛЕКСАНДРОВИЧ=13, ВИКТОРОВИЧ=10...):
        # верхняя граница 16, иначе отчества длиннее 12 букв не достижимы
        for length in range(7, min(len(t) - 3, 16)):
            head = t[:-length]
            stem = t[-length:][:-4]
            if len(stem) < 3:
                continue
            # ствол отчества — точно имя из словаря? тогда голова
            # (фамилия ИЛИ имя, напр. АЛЕКСАНДР+МИХАЙЛОВИЧ) не обязана
            # оканчиваться ОВ/ЕВ — принимаем
            if stem not in _FIRST_NAMES and not _ok(head):
                continue
            score = (0 if stem in _FIRST_NAMES else 1, length)
            if best is None or score < best[0]:
                best = (score, length)
        if best:
            length = best[1]
            parts.append(t[-length:])
            t = t[:-length]
            # голова — само по себе имя (имя+отчество без фамилии)
            if t in _FIRST_NAMES:
                return [t] + parts
    L = len(t)
    hi = min(L - 4, 14)
    for k in range(hi, 2, -1):
        if t[L - k:] in _FIRST_NAMES and L - k >= 4:
            head = t[:L - k]
            if len(head) <= 15 and head not in _FIRST_NAMES:
                return [head, t[L - k:]] + parts
    if parts or not tok.endswith(_PAT_ENDS):
        for i in range(5, L - 3):
            name = t[i:]
            if not t[:i].endswith(("ОВ", "ЕВ")):
                continue
            if name in _FIRST_NAMES:
                return [t[:i], name] + parts
            if len(name) >= 5 and not name.endswith(
                    ("СКИЙ", "СКАЯ", "СКОГО", "СКОЙ", "СКИЕ")):
                return [t[:i], name] + parts
    return None


def _norm_name(s):
    s = re.sub(r"[^\u0400-\u04FF]", "", s or "").lower()
    return s


_END_NAME = (
    r"(?:ОКПО\s*\d+|ОКПО\s*\d+|ОГРНИП\s*\d+|ИНН\s*\d+|ОКПО\s*\d+|ОГРНИП\s*\d+|"
    r"Место\s+рождения|место\s+рождения|Дата\s+рождения|дата\s+рождения|,\s|\.\s)")

_NAME_PATTERNS = [
    # Извещение: «Адресат: индивидуальный предприниматель ФИО ОКПО/ИНН/Место рождения...»
    ("fiz", re.compile(
        r"Адресат\s*[:,]\s*"
        r"([A-Za-z\u0400-\u04FFа-яё][\u0400-\u04FFа-яёA-Za-z \-]{0,95}?)\s*"
        + _END_NAME)),
    # Протокол: «...свершающий (свершавший): ... ФИО ОКПО...» / «нарушитель-организация: ФИО ОКПО»
    ("fiz", re.compile(
        r"(?:соверш\w+\s*\(.*?\)\s*[:,]|нарушитель\w*[-\s]{0,3}организация\w*\s*[:,])\s*"
        r"([A-Za-z\u0400-\u04FFа-яё][\u0400-\u04FFа-яёA-Za-z \-]{0,95}?)\s*"
        + _END_NAME)),
    # Определение: «в отношении ... ФИО ОКПО...»
    ("fiz", re.compile(
        r"в\s+отношении\s+(?:\w+\s+)*?"
        r"([A-Za-z\u0400-\u04FFа-яё][\u0400-\u04FFа-яёA-Za-z \-]{0,95}?)\s*"
        + _END_NAME)),
    # Подтверждение вручения: «Организация нарушителя: ФИО.»
    ("fiz", re.compile(
        r"Организация\s+(?:нарушителя|правонарушителя)\s*[:,]\s*"
        r"([A-Za-z\u0400-\u04FF][\u0400-\u04FFа-яёA-Za-z \-]{0,95}?)\s*[.;:,]")),
    # ЮЛ в кавычках: «ООО «Название» ...»
    ("yur", re.compile(r"(?:ООО|АО|ЗАО|ПАО|универсальное предприятие|корпорация)\s*[«“]([^»”\"]{2,100})[»”\"]")),
    # ЮЛ без кавычек: «ООО Название ... ОКПО/ИНН»
    ("yur", re.compile(r"(?:ООО|ЗАО|ПАО)\s+([A-Za-z\u0400-\u04FFа-яё][\u0400-\u04FFа-яёA-Za-z \-]{0,95}?)\s+(?:ИНН\s*\d+|ОКПО\s*\d+|зарегистрир\w*)")),
]


def _unknown_toks(toks):
    """Число токенов, которые НЕ имя и НЕ отчество (фамилия допустима — 1).
    99 = кандидат отбрасывается (есть ОКПО/ИНН/слишком длинное)."""
    if len(toks) < 2:
        return 99
    n = 0
    for t in toks:
        if (t in ("ОКПО", "ИНН", "КПП", "ОГРНИП", "ОГРН", "АДРЕС", "АДРЕСА",
                  "БУХГАЛТЕР", "ИНДИВИДУАЛЬНЫЙ", "ОБЩЕСТВА") or len(t) > 20):
            return 99
        if t.upper() in _FIRST_NAMES:
            continue
        if t.upper().endswith(("ОВИЧ", "ЕВИЧ", "ОВНА", "ЕВНА", "ИЧНА", "ИВНА")):
            continue
        if 4 <= len(t) <= 13:
            n += 1
    return n


def _body_fix(flat, name):
    """Если в ФИО есть подозрительный склеенный кусок — ищем чистый вариант
    'предприниматель + 2-4 слова' в теле документа. Возвращает строку или None."""
    best = None
    best_score = _unknown_toks(name.split())
    for m in re.finditer(
            r"предприниматель\s+([А-ЯЁа-яё]{3,}(?:\s+[А-ЯЁа-яё]{2,}){1,3})", flat):
        toks = m.group(1).split()
        good = []
        for t in toks:
            if (t in ("ОКПО", "ИНН", "КПП", "ОГРНИП", "ОГРН", "АДРЕС", "АДРЕСА",
                      "БУХГАЛТЕР") or len(t) > 20):
                break
            good.append(t)
        if len(good) < 2:
            continue
        score = _unknown_toks(good)
        if score < best_score:
            best, best_score = " ".join(good).upper(), score
    return best


def extract_recipient(raw_text, issuer_inn):
    """Вытаскивает из текста PDF данные получателя: тип, ФИО/название, ИНН, адрес."""
    t = _flat_text(raw_text)
    rec = {"kind": None, "name": "", "org": "", "inn": None, "address": None}

    for kind, pat in _NAME_PATTERNS:
        m = pat.search(t)
        if not m:
            continue
        toks = _clean_name_tokens(m.group(1))
        if not toks or sum(len(x) for x in toks) < 5:
            continue
        if kind == "yur":
            rec["kind"] = "yur"
            rec["org"] = " ".join(toks)
            rec["name"] = rec["org"]
        else:
            rec["kind"] = "fiz"
            rec["name"] = " ".join(toks)
        break

    # ИНН (приоритет — закреплённые за получателем: «ОКПО ... ИНН ...»)
    cands = []
    for mm in re.finditer(r"ОКПО\s*\d+\s*ИНН\s*(\d{10,12})", t):
        cands.append((0, mm.group(1)))
    for mm in re.finditer(r"ОКПО\s*\d+\s*ОКПО\s*(\d{10,12})", t):
        cands.append((0, mm.group(1)))
    for mm in re.finditer(r"ИНН\s*(\d{12})\s*ОГРНИП\s*\d+", t):
        cands.append((1, mm.group(1)))
    for mm in re.finditer(r"ИНН\s*(\d{10})\s*КПП\s*\d+", t):
        cands.append((2, mm.group(1)))
    for mm in re.finditer(r"ИНН\s*(\d{10,12})", t):
        cands.append((9, mm.group(1)))
    cands.sort(key=lambda x: x[0])
    for _, v in cands:
        if v and v != issuer_inn:
            rec["inn"] = v
            break

    rec["address"] = extract_address(t)
    if rec["name"]:
        fixed = _body_fix(t, rec["name"])
        if fixed:
            rec["name"] = fixed
    return rec


_ADDR_START = re.compile(
    r"адреса?\s*(?:места\s*)?(?:регистрации|жительства)|зарегистрирован\w*\s+по\s+адресу|адреса?\s+получателя|адрес\w*\s*:",
    re.I)
# стоп-слова БЕЗ \b — текст PDF может быть без пробелов ("...Д. 4ИНН 321...");
# ИНН — только когда дальше НЕ буква (иначе "ИННОВАЦИЯ" режет адрес)
_ADDR_STOP = re.compile(
    r"\s*(?:тел|телефон|ОКПО|ИНН(?![а-яёА-ЯЁ])|ОГРНИП|ОГРН|паспорт\s|место\s+рождения|дата\s+рождения|далее\s*–?\s*лицо|\(далее|занимаем\w*\s+должност\w*|должност\w*\s*[:—-]|Рассмотрен\w*\s+документ|Докладчик|Основание|Рассмотрев\s|номер|телефона|\d{1,2}\s+(?:январ|феврал|март|апрел|ма[е][а]?|июн[е]?|июл[е]?|август|сентябр|октябр|ноябр|декабр)\w*|вызывает|предприниматель\s)",
    re.I)


_DATE_CUT = re.compile(
    r"\d{1,2}\s+(?:январ|феврал|март|апрел|ма[е][а]?|июн[е]?|июл[е]?|август|сентябр|октябр|ноябр|декабр)\w*",
    re.I)


def _street_date_at(s, pos):
    """Дата в pos — часть названия улицы ("УЛ. 23 СЕНТЯБРЯ, Д. 2")? Тогда это не стоп-слово."""
    dm = _DATE_CUT.match(s, pos)
    if not dm:
        return False
    before = s[max(0, pos - 40):pos].rstrip(" \t,.\u00a0\u2014-—")
    after = s[dm.end():dm.end() + 30]
    return (bool(re.search(
                r"\b(?:УЛ|УЛИЦА|ПЕР|ПРОСП|Б-Р|Б-РА|БУЛЬВАР|ПР[ОЕ]Д|ШОССЕ|ПЛОЩ)\b[^\d]{0,5}$",
                before, re.I))
            and bool(re.match(r"[^\d]{0,3}(?:Д|ДОМ)\b[^\d]{0,3}\d+", after, re.I)))
# концовка адреса: Д./Дом + н. (+ опц. К/Кор/Корп/Кв + н. или буква)
_ADDR_END = re.compile(
    r"(?<![\u0400-\u044fA-Za-z])(?:КВ|КВАРТИРА|КОРП|КОР|ЛИТ|ЛИТЕРА|Д|ДОМ|К)[\s.,]*[0-9А-ЯЁ]{1,3}\b",
    re.I)


def extract_address(t):
    """Адрес регистрации/проживания до стоп-слоva (телефон/ОКПО/...).
    Кончается на кв./квартиру или на Д./Дом (+кorpus), дата и хворост режутся."""
    best = None
    for m in _ADDR_START.finditer(t):
        seg = t[m.end():m.end() + 320]
        st = None
        for stc in _ADDR_STOP.finditer(seg):
            if _street_date_at(seg, stc.start()):
                continue  # "УЛ. 23 СЕНТЯБРЯ, Д. 2, КВ. 14" — дата — название улицы
            st = stc
            break
        a = seg[:st.start()] if st else seg
        a = re.sub(r"\s+", " ", a)
        dm = None
        for dmc in _DATE_CUT.finditer(a):
            if _street_date_at(a, dmc.start()):
                continue
            dm = dmc
            break
        if dm:
            a = a[:dm.start()]
        ems = list(_ADDR_END.finditer(a))
        if ems:
            a = a[:ems[-1].end()]
        a = a.strip(" \t.,;:«»\u00ab\u00bb\"'“”\u201e()[]")
        ok = re.match(r"^\d{6}[\s,]", a) or len(re.findall(
            r"\b(?:обл|область|район|город|города|г\.|село|деревня|ул|пер|поселение|кв|д)\b", a, re.I)) >= 2
        if len(a) < 8 or not ok:
            continue
        score = (1 if re.match(r"^\d{6}\s", a) else 0, len(a))
        if best is None or score > best[0]:
            best = (score, a)
    return best[1] if best else None


def _trim_co_name_prefix(a):
    """Адрес из PDF иногда начинается с полного наименования фирмы
    («ОБЩЕСТВО СОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ ...«) + «, почтовый код, адрес».
    Обрезать префикс наименования, если после первой запятой есть почтовый код."""
    if not a:
        return a
    if not re.search(r"^\s*(?:ОБЩЕСТВО\s*СО[ГЩ]РАНИЧЕННОЙ\s+ОТВЕТСТВЕННОСТЬЮ|АКЦИОНЕРНОЕ\s+ОБЩЕСТВО)", a, re.I):
        return a
    if "," in a:
        rest = a.split(",", 1)[1].strip()
        if re.search(r"\d{6}", rest):
            return rest
    return a


def parse_xml_recipient(xml_path):
    """Данные получателя из XML-формы LKS (надёжнее, чем парсинг PDF).

    Фактическая покрываемость (проверено по 253 формам за 17.03):
      • participantType  — 253/253 (IP или UL)
      • subjectObservationObjectShortName — все UL (наименование организации)
      • CaseSubjectInn   — 58/70 UL (в JSON parentDocumentList; 12 UL — нет)
      • offenceLocation  — 183/183 IP (адрес прописки) + 61/70 UL (адрес юр.лица)
      • ФИО/ИНН отдельных граждан (IP) — отсутствуют, только в PDF
    Возвращает dict: ptype 'UL'|'IP'|None, org, inn, addr (где нет — None).
    """
    res = {"ptype": None, "org": None, "inn": None, "addr": None}
    try:
        with open(xml_path, "rb") as f:
            raw = f.read().decode("utf-8", "replace")
    except Exception:
        return res
    raw = re.sub(r">\s+<", "><", raw)

    def g(key):
        m = re.search(
            r"<%s>\s*<id>[^<]*</id>\s*<value>(.*?)</value>" % re.escape(key),
            raw, re.S)
        return m.group(1).strip() if m else None

    pt = g("offenceParticipantType")
    if pt in ("IP", "UL"):
        res["ptype"] = pt
    org = g("subjectObservationObjectShortName")
    if org and len(org) >= 3:
        res["org"] = org
    # CaseSubjectInn — НЕ отдельный XML-элемент, а ключ внутри JSON в parentDocumentList
    m = re.search(r'CaseSubjectInn["\s:]*"?(\d{10})', raw)
    if m:
        res["inn"] = m.group(1)
    loc = g("offenceLocation")
    if loc and re.search(r"\d{6}", loc):  # адрес — только с почтовым индексом
        res["addr"] = re.sub(r"\s+", " ", loc).strip().rstrip(" ,;:.")
    return res


def build_registry_xlsx(template, rpath, subset, sender_address):
    """Заполняет один реестр (шаблон -> rpath) по списку групп subset."""
    import openpyxl
    shutil.copyfile(template, rpath)
    wb = openpyxl.load_workbook(rpath)
    ws = None
    for sn in wb.sheetnames:
        if "реестр" in sn.lower():
            ws = wb[sn]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]
    row = 2
    for g in subset:
        ws.cell(row=row, column=1, value=g["zipname"])
        ws.cell(row=row, column=2, value=1)   # тип письма: заказное
        ws.cell(row=row, column=3, value=1 if g["kind"] == "yur" else 0)
        if g["kind"] == "yur" and g["best_org"]:
            ws.cell(row=row, column=4, value=g["best_org"])
        if g["inn"]:
            ws.cell(row=row, column=5,
                    value=int(g["inn"]) if g["inn"].isdigit() else g["inn"])
        # столбец F (КПП) — по требованию не заполняется
        if g["kind"] != "yur" and g["best_name"]:
            tok = g["best_name"].split()
            if len(tok) >= 3:
                fam, im, otc = tok[0], tok[1], " ".join(tok[2:])
            elif len(tok) == 2:
                fam, im, otc = tok[0], tok[1], ""
            else:
                fam, im, otc = tok[0], "", ""
            ws.cell(row=row, column=7, value=fam)
            if im:
                ws.cell(row=row, column=8, value=im)
            if otc:
                ws.cell(row=row, column=9, value=otc)
        if g["best_addr"]:
            ws.cell(row=row, column=10, value=g["best_addr"])
        if sender_address:
            ws.cell(row=row, column=11, value=sender_address)
        # столбец L (доп. услуга ЮЗЭУВ) — по требованию не заполняется
        row += 1
    wb.save(rpath)
    return row - 2


def make_post_zips(date_dir, groups, cfg):
    """Собирает post.zip (<=50 писем) или post1.zip/post2.zip/... при >50.

    Каждый архив: registry.xlsx (только его получатели) + сжатые PDF.
    Возвращает список путей архивов.
    """
    template = os.path.join(BASE_DIR, "registry-template.xlsx")
    if not os.path.exists(template):
        print("ОШИБКА: не найден шаблон реестра: %s" % template)
        return []
    sender_address = (cfg.get("sender_address") or "").strip()
    if not sender_address:
        sender_address = "241030, БРЯНСКАЯ обл., г. БРЯНСК, ул. КРАСНОАРМЕЙСКАЯ, д. 60"
        print("  (адрес отправителя по умолчанию — задайте sender_address в config.json)")

    groups_sorted = sorted(groups, key=lambda g: (g["inn"] is None, g["inn"] or "",
                                                  (g["best_org"] or g["best_name"]).lower()))
    total = len(groups_sorted)
    chunks = [groups_sorted[i:i + ZIP_MAX_DOCS] for i in range(0, total, ZIP_MAX_DOCS)]
    multi = len(chunks) > 1
    zips = []
    for idx, chunk in enumerate(chunks, 1):
        zname = ("post%d.zip" % idx) if multi else "post.zip"
        rname = ("registry_%d.xlsx" % idx) if multi else "registry.xlsx"
        registry_path = os.path.join(date_dir, rname)
        n_rows = build_registry_xlsx(template, registry_path, chunk, sender_address)
        print("  РЕЕСТР под %s: %d строк" % (zname, n_rows))
        zip_path = os.path.join(date_dir, zname)
        if os.path.exists(zip_path):
            os.remove(zip_path)
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(registry_path, arcname="registry.xlsx")
            for g in chunk:
                z.write(g["merged"], arcname=g["zipname"])
        zips.append(zip_path)
    return zips


def stage_post(date_dir, cfg):
    """Стадия 2: парсит получателя, сливает PDF на адресата, строит реестр(ы) и post*.zip."""
    try:
        import pdfplumber
        from pypdf import PdfWriter
        import openpyxl
    except ImportError:
        print("\nОШИБКА: для почтовой стадии нужны pypdf, pdfplumber, openpyxl, pikepdf.")
        print("Установка:  pip install pypdf pdfplumber openpyxl pikepdf")
        sys.exit(1)

    print()
    print("=" * 62)
    print("Стадия 2: почтовый пакет (Почта РФ)")
    print("=" * 62)

    issuer_inn = (cfg.get("login") or "").split("_")[0]

    # --- собираем исходные PDF (только каталоги типов)
    docs = []
    for tname in TYPE_ORDER:
        tdir = os.path.join(date_dir, tname)
        if not os.path.isdir(tdir):
            continue
        for f in sorted(os.listdir(tdir)):
            if f.lower().endswith(".pdf") and not f.startswith("."):
                docs.append({"file": os.path.join(tdir, f), "type": tname, "rec": None})
    if not docs:
        print("Почтовая стадия: нет исходных PDF в %s — пропускаю." % date_dir)
        return
    print("Документов к обработке: %d" % len(docs))

    # --- нормализация: Ё -> Е в ФИО (Почта РФ отклоняет букву Ё)
    for rec in docs:
        k = _normalize_yo_pdf(rec["file"])
        if k:
            print("  Ё->Е: %s (%d знака)" % (os.path.basename(rec["file"]), k))

    # --- извлекаем получателя из каждого PDF
    for rec in docs:
        try:
            with pdfplumber.open(rec["file"]) as pdf:
                raw = "\n".join((p.extract_text() or "") for p in pdf.pages)
        except Exception as e:
            print("  [НЕ СМОГ прочесть] %s: %s" % (os.path.basename(rec["file"]), e))
            continue
        why = _is_unwanted_text(_flat_text(raw))
        if why:
            print("  [пропущено] %s — %s" % (os.path.basename(rec["file"]), why))
            continue
        rec["rec"] = extract_recipient(raw, issuer_inn)
        r = rec["rec"]
        # --- XML-форма LKS: авторитетный источник для типа получателя,
        #     наименования/ИНН юр.лиц и адреса (PDF — резерв)
        xpath = os.path.splitext(rec["file"])[0] + ".xml"
        if os.path.exists(xpath):
            xr = parse_xml_recipient(xpath)
            if xr["ptype"] == "UL":
                r["kind"] = "yur"
                if xr["org"]:
                    r["org"] = xr["org"]
                    r["name"] = xr["org"]
                if xr["inn"]:
                    r["inn"] = xr["inn"]
                if xr["addr"]:
                    r["address"] = xr["addr"]
            elif xr["ptype"] == "IP" and xr["addr"]:
                r["address"] = xr["addr"]
        r["address"] = _trim_co_name_prefix(r["address"])
        print("  %s/%s: %s | ИНН=%s | адрес=%s" % (
            rec["type"], os.path.splitext(os.path.basename(rec["file"]))[0],
            r["org"] or r["name"] or "(не определён)",
            r["inn"] or "-", "есть" if r["address"] else "НЕТ"))

    with_info = [x for x in docs if x["rec"] and (x["rec"]["name"] or x["rec"]["org"] or x["rec"]["inn"])]
    if not with_info:
        print("ОШИБКА: не удалось распознать ни одного получателя.")
        return

    # --- группируем: ключ — ИНН, резерво — нормализованное название/ФИО
    groups, by_inn, by_name = [], {}, {}
    for rec in with_info:
        r = rec["rec"]
        nk = _norm_name(r["org"] or r["name"])
        g = None
        if r["inn"] and r["inn"] in by_inn:
            g = by_inn[r["inn"]]
        elif nk and nk in by_name:
            g = by_name[nk]
        if g is None:
            g = {"inn": r["inn"], "kind": r["kind"] or "fiz", "names": [], "orgs": [],
                 "addresses": [], "docs": [], "zipname": None, "merged": None}
            groups.append(g)
        g["docs"].append(rec)
        if r["kind"] == "yur":
            g["kind"] = "yur"
        if r["name"]:
            g["names"].append(r["name"])
        if r["org"]:
            g["orgs"].append(r["org"])
        if r["inn"] and not g["inn"]:
            g["inn"] = r["inn"]
        if r["address"]:
            g["addresses"].append(r["address"])
        if r["inn"]:
            by_inn[r["inn"]] = g
        if nk:
            by_name[nk] = g

    for g in groups:
        g["best_name"] = max(g["names"], key=len) if g["names"] else ""
        g["best_org"] = max(g["orgs"], key=len) if g["orgs"] else ""
        g["best_addr"] = (max(g["addresses"],
                              key=lambda a: (1 if re.match(r"^\d{6}\s", a) else 0, len(a)))
                          if g["addresses"] else "")
        g["docs"].sort(key=lambda r: (TYPE_ORDER.get(r["type"], 9), r["file"]))
        if not g["best_name"]:
            print("  ВНИМАНИЕ: у группы нет ФИО: %s" % [
                os.path.basename(x["file"]) for x in g["docs"]])
        if not g["best_addr"]:
            print("  ВНИМАНИЕ: для группы %s адрес не найден — заполните вручную в реестре" %
                  (g["best_name"] or g["inn"]))

    # --- слияние PDF в один файл на получателя
    merged_dir = os.path.join(date_dir, "merged")
    shutil.rmtree(merged_dir, ignore_errors=True)
    os.makedirs(merged_dir, exist_ok=True)
    used = set()
    for g in groups:
        if g["inn"] and g["inn"].isdigit():
            base = g["inn"]
        else:
            label = (g["best_org"] if g["kind"] == "yur" else g["best_name"]) or "получатель"
            base = re.sub(r"[\\/:*?\"<>|]+", "_", label).strip(" ._")
            if g["inn"]:
                base += "_ИНН" + g["inn"]
        base = re.sub(r"\s+", " ", base)[:230].strip(" .")
        zname = base + ".pdf"
        k = 1
        while zname.lower() in used:
            zname = "%s_%d.pdf" % (base, k); k += 1
        used.add(zname.lower())
        g["zipname"] = zname
        dest = os.path.join(merged_dir, zname)
        if len(g["docs"]) == 1:
            shutil.copyfile(g["docs"][0]["file"], dest)
        else:
            w = PdfWriter()
            for r in g["docs"]:
                w.append(r["file"])
            with open(dest, "wb") as fo:
                w.write(fo)
        g["merged"] = dest
        print("  СЛИТО: %s  (%d док.: %s)" % (
            zname, len(g["docs"]),
            ", ".join(os.path.basename(r["file"]) for r in g["docs"])))

        # --- реестр + архив(ы) с разбиением по ZIP_MAX_DOCS писем
    zips = make_post_zips(date_dir, groups, cfg)
    if zips:
        for zp in zips:
            print("  %s (%d КБ)" % (zp, os.path.getsize(zp) // 1024))
        print("ИТОГО: %d получателей, архив(ов): %d" % (len(groups), len(zips)))


def main():
    flags = [a for a in sys.argv[1:] if a.startswith("--")]
    pos = [a for a in sys.argv[1:] if not a.startswith("--")]
    post_only = "--post-only" in flags
    no_post = "--no-post" in flags
    cfg = load_config()
    d = ask_date(pos[0] if pos else None)
    date_ddmmyyyy = d.strftime("%d.%m.%Y")
    date_iso = d.strftime("%Y-%m-%d")
    date_dir = os.path.join(OUT_ROOT, date_ddmmyyyy)
    print("Дата утверждения: %s (каталог: %s)" % (date_ddmmyyyy, date_dir))

    rc = 0
    if not post_only:
        rc = do_download(cfg, d, date_iso, date_ddmmyyyy)
    if not no_post:
        stage_post(date_dir, cfg)
        cleanup_date_dir(date_dir)
    return rc


def cleanup_date_dir(date_dir):
    """После успешной сборки: в out/<дата>/ остаются только post*.zip."""
    import shutil
    zips = sorted(
        e for e in os.listdir(date_dir)
        if re.match(r"^post\d?\.zip$", e)
        and os.path.isfile(os.path.join(date_dir, e))
        and os.path.getsize(os.path.join(date_dir, e)) > 0)
    if not zips:
        print("ОЧИСТКА ПРОПУЩЕНА: post.zip не найден/пуст — исходные файлы сохранены.")
        return
    removed = 0
    for e in os.listdir(date_dir):
        if re.match(r"^post\d?\.zip$", e):
            continue
        full = os.path.join(date_dir, e)
        try:
            if os.path.isdir(full) and not os.path.islink(full):
                shutil.rmtree(full)
            else:
                os.remove(full)
            removed += 1
        except OSError as ex:
            print("  не удалено: %s (%s)" % (e, ex))
    if removed:
        print("ОЧИСТКА: в %s оставлены только %s (удалено объектов: %d)" %
              (date_dir, ", ".join(zips), removed))


if __name__ == "__main__":
    sys.exit(main())
